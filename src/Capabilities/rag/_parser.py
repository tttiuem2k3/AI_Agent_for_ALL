# -*- coding: utf-8 -*-
"""RAG parsers and adapters for supported document formats."""
from __future__ import annotations

import asyncio
import io
import json
import mimetypes
from abc import ABC, abstractmethod
from pathlib import Path
from typing import Any, Literal

from Capabilities.document_conversion import ConversionOptions, DocumentConversionService
from Runtime.message import TextBlock

from ._document import Section


class ParserBase(ABC):
    """Abstract parser for one file family."""

    supported_media_types: list[str]

    @classmethod
    def supported_extensions(cls) -> list[str]:
        extensions: set[str] = set()
        for media_type in cls.supported_media_types:
            extensions.update(mimetypes.guess_all_extensions(media_type))
        return sorted(extensions)

    @abstractmethod
    async def parse(self, file: bytes | str, filename: str) -> list[Section]:
        """Parse bytes or a file path into sections."""


class WordParser(ParserBase):
    """Parse ``.docx`` files into text sections."""

    supported_media_types = [
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ]

    @classmethod
    def supported_extensions(cls) -> list[str]:
        return [".docx"]

    async def parse(self, file: bytes | str, filename: str) -> list[Section]:
        try:
            from docx import Document
        except ImportError as exc:
            raise ImportError(
                "WordParser requires optional dependency `python-docx`. "
                "Install with `asoft-ai-services[rag]`.",
            ) from exc

        document = Document(file if isinstance(file, str) else io.BytesIO(file))
        parts: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                parts.append(text)
        for table in document.tables:
            rows = [
                [cell.text.strip() for cell in row.cells]
                for row in table.rows
            ]
            if rows:
                parts.append(_table_to_markdown(rows))
        if not parts:
            return []
        return [
            Section(
                content=TextBlock(text="\n\n".join(parts)),
                source=filename,
                metadata={},
            ),
        ]


class ExcelParser(ParserBase):
    """Parse Excel files into one section per sheet or a merged section."""

    supported_media_types = [
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel",
    ]

    def __init__(
        self,
        *,
        separate_sheet: bool = False,
        table_format: Literal["markdown", "json"] = "markdown",
    ) -> None:
        self.separate_sheet = separate_sheet
        self.table_format = table_format

    @classmethod
    def supported_extensions(cls) -> list[str]:
        return [".xls", ".xlsx"]

    async def parse(self, file: bytes | str, filename: str) -> list[Section]:
        try:
            import pandas as pd
        except ImportError as exc:
            raise ImportError(
                "ExcelParser requires optional dependencies `pandas` and "
                "`openpyxl`/`xlrd`. Install with `asoft-ai-services[rag]`.",
            ) from exc

        excel_file = pd.ExcelFile(file if isinstance(file, str) else io.BytesIO(file))
        sections: list[Section] = []
        try:
            for sheet_name in excel_file.sheet_names:
                frame = excel_file.parse(sheet_name=sheet_name)
                if frame.empty:
                    continue
                rows = _dataframe_to_rows(frame)
                text = (
                    _table_to_markdown(rows)
                    if self.table_format == "markdown"
                    else json.dumps(rows, ensure_ascii=False)
                )
                sections.append(
                    Section(
                        content=TextBlock(text=text),
                        source=filename,
                        metadata={"sheet": sheet_name},
                    ),
                )
        finally:
            excel_file.close()

        if self.separate_sheet or not sections:
            return sections
        return [
            Section(
                content=TextBlock(
                    text="\n\n".join(section.content.text for section in sections),
                ),
                source=filename,
                metadata={},
            ),
        ]


def _dataframe_to_rows(frame: Any) -> list[list[str]]:
    headers = [str(column) for column in frame.columns]
    rows = [headers]
    for values in frame.fillna("").astype(str).values.tolist():
        rows.append([str(value) for value in values])
    return rows


def _table_to_markdown(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    header = normalized[0]
    separator = ["---"] * width
    body = normalized[1:]
    all_rows = [header, separator, *body]
    return "\n".join("| " + " | ".join(row) + " |" for row in all_rows)


class DocumentConversionParser(ParserBase):
    """RAG adapter backed by the generic document conversion capability."""

    supported_media_types: list[str] = []

    def __init__(self, converter: DocumentConversionService | None = None) -> None:
        self.converter = converter or DocumentConversionService()

    @classmethod
    def supported_extensions(cls) -> list[str]:
        return [
            ".csv", ".doc", ".docm", ".docx", ".epub", ".odp", ".ods",
            ".odt", ".pdf", ".pot", ".pps", ".ppsm", ".ppsx", ".ppt",
            ".pptm", ".pptx", ".rtf", ".xls", ".xlsb", ".xlsm", ".xlsx",
        ]

    async def parse(self, file: bytes | str, filename: str) -> list[Section]:
        content = file if isinstance(file, bytes) else await asyncio.to_thread(Path(file).read_bytes)
        converted = await self.converter.convert(
            content,
            filename=filename,
            options=ConversionOptions(include_structure=False, include_assets=False),
        )
        if not converted.markdown.strip():
            return []
        return [
            Section(
                content=TextBlock(text=converted.markdown),
                source=filename,
                metadata={
                    "format": converted.format.value,
                    "warnings": list(converted.warnings),
                },
            ),
        ]
